"""
AI Literature Review System - Modern Academic UI
=================================================
Clean, minimal, professional interface for automated literature review.

Design Principles:
- Off-white background (#F7F8FA)
- White cards with subtle shadows
- Soft blue accent (#3B82F6)
- Inter/system-ui fonts
- Generous whitespace
- Academic, presentation-ready aesthetic
"""

import os
import sys
import json
import time
import logging
import tempfile
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple
from pathlib import Path

# Ensure project root on path
ROOT = Path(__file__).parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

os.chdir(ROOT)

# Setup logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

import gradio as gr

# Import modules
from modules.search_papers import search_papers, rank_papers, save_metadata
from modules.download_pdf import download_papers, save_download_metadata
from modules.extract_text import extract_text_for_papers
from modules.analyze_text import process_analysis
from modules.generate_draft import generate_drafts
from modules.critique_draft import critique_drafts
from urllib.parse import urlparse

# =============================================================================
# CUSTOM CSS - Academic Theme
# =============================================================================

CUSTOM_CSS = """
/* ===== ROOT VARIABLES ===== */
:root {
    --bg-primary: #F7F8FA;
    --bg-card: #FFFFFF;
    --accent-primary: #3B82F6;
    --accent-secondary: #6366F1;
    --accent-success: #10B981;
    --accent-warning: #F59E0B;
    --text-primary: #1F2937;
    --text-secondary: #6B7280;
    --text-muted: #9CA3AF;
    --border-light: #E5E7EB;
    --shadow-sm: 0 1px 2px rgba(0, 0, 0, 0.05);
    --shadow-md: 0 4px 6px -1px rgba(0, 0, 0, 0.1), 0 2px 4px -1px rgba(0, 0, 0, 0.06);
    --shadow-lg: 0 10px 15px -3px rgba(0, 0, 0, 0.1), 0 4px 6px -2px rgba(0, 0, 0, 0.05);
    --radius-sm: 8px;
    --radius-md: 12px;
    --radius-lg: 16px;
}

/* ===== GLOBAL STYLES ===== */
.gradio-container {
    background: var(--bg-primary) !important;
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif !important;
    max-width: 1200px !important;
    margin: 0 auto !important;
}

/* ===== HEADER STYLES ===== */
.header-container {
    background: linear-gradient(135deg, var(--accent-primary) 0%, var(--accent-secondary) 100%);
    color: white;
    padding: 2rem 2.5rem;
    border-radius: var(--radius-lg);
    margin-bottom: 1.5rem;
    box-shadow: var(--shadow-lg);
}

.header-title {
    font-size: 2rem;
    font-weight: 700;
    margin: 0 0 0.5rem 0;
    letter-spacing: -0.025em;
}

.header-subtitle {
    font-size: 1rem;
    opacity: 0.9;
    font-weight: 400;
}

/* ===== CARD STYLES ===== */
.card {
    background: var(--bg-card) !important;
    border-radius: var(--radius-md) !important;
    box-shadow: var(--shadow-md) !important;
    border: 1px solid var(--border-light) !important;
    padding: 1.5rem !important;
    margin-bottom: 1rem !important;
}

.card-title {
    font-size: 0.875rem;
    font-weight: 600;
    color: var(--text-secondary);
    text-transform: uppercase;
    letter-spacing: 0.05em;
    margin-bottom: 1rem;
}

/* ===== INPUT STYLES ===== */
.input-container input,
.input-container textarea {
    border: 2px solid var(--border-light) !important;
    border-radius: var(--radius-sm) !important;
    padding: 0.75rem 1rem !important;
    font-size: 1rem !important;
    transition: border-color 0.2s ease, box-shadow 0.2s ease !important;
}

.input-container input:focus,
.input-container textarea:focus {
    border-color: var(--accent-primary) !important;
    box-shadow: 0 0 0 3px rgba(59, 130, 246, 0.1) !important;
    outline: none !important;
}

/* ===== BUTTON STYLES ===== */
.primary-button {
    background: linear-gradient(135deg, var(--accent-primary) 0%, var(--accent-secondary) 100%) !important;
    color: white !important;
    border: none !important;
    border-radius: var(--radius-sm) !important;
    padding: 0.875rem 2rem !important;
    font-weight: 600 !important;
    font-size: 1rem !important;
    cursor: pointer !important;
    transition: transform 0.2s ease, box-shadow 0.2s ease !important;
    box-shadow: var(--shadow-md) !important;
}

.primary-button:hover {
    transform: translateY(-1px) !important;
    box-shadow: var(--shadow-lg) !important;
}

.secondary-button {
    background: var(--bg-card) !important;
    color: var(--text-primary) !important;
    border: 2px solid var(--border-light) !important;
    border-radius: var(--radius-sm) !important;
    padding: 0.75rem 1.5rem !important;
    font-weight: 500 !important;
    cursor: pointer !important;
    transition: all 0.2s ease !important;
}

.secondary-button:hover {
    border-color: var(--accent-primary) !important;
    color: var(--accent-primary) !important;
}

/* ===== PROGRESS STEPPER ===== */
.progress-stepper {
    display: flex;
    justify-content: space-between;
    align-items: center;
    padding: 1rem 0;
    position: relative;
}

.step {
    display: flex;
    flex-direction: column;
    align-items: center;
    gap: 0.5rem;
    flex: 1;
    position: relative;
    z-index: 1;
}

.step-icon {
    width: 32px;
    height: 32px;
    border-radius: 50%;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 0.875rem;
    font-weight: 600;
    transition: all 0.3s ease;
}

.step-pending .step-icon {
    background: var(--border-light);
    color: var(--text-muted);
}

.step-active .step-icon {
    background: var(--accent-primary);
    color: white;
    animation: pulse 1.5s infinite;
}

.step-complete .step-icon {
    background: var(--accent-success);
    color: white;
}

.step-label {
    font-size: 0.75rem;
    color: var(--text-secondary);
    font-weight: 500;
    text-align: center;
}

@keyframes pulse {
    0%, 100% { transform: scale(1); }
    50% { transform: scale(1.1); }
}

/* ===== TABS ===== */
.tabs-container .tab-nav {
    background: var(--bg-primary) !important;
    border-radius: var(--radius-sm) !important;
    padding: 4px !important;
    gap: 4px !important;
}

.tabs-container .tab-nav button {
    border-radius: var(--radius-sm) !important;
    padding: 0.75rem 1.25rem !important;
    font-weight: 500 !important;
    transition: all 0.2s ease !important;
    border: none !important;
    background: transparent !important;
    color: var(--text-secondary) !important;
}

.tabs-container .tab-nav button.selected {
    background: var(--bg-card) !important;
    color: var(--accent-primary) !important;
    box-shadow: var(--shadow-sm) !important;
}

/* ===== CONTENT AREA ===== */
.content-area {
    background: var(--bg-card);
    border-radius: var(--radius-md);
    padding: 1.5rem;
    min-height: 300px;
    max-height: 500px;
    overflow-y: auto;
    border: 1px solid var(--border-light);
    line-height: 1.7;
    font-size: 0.9375rem;
    color: var(--text-primary);
}

.content-area h1, .content-area h2, .content-area h3 {
    color: var(--text-primary);
    margin-top: 1.5rem;
    margin-bottom: 0.75rem;
}

.content-area h2 {
    font-size: 1.25rem;
    border-bottom: 2px solid var(--border-light);
    padding-bottom: 0.5rem;
}

/* ===== ACTION BAR ===== */
.action-bar {
    display: flex;
    gap: 0.75rem;
    justify-content: flex-end;
    padding: 1rem 0;
    border-top: 1px solid var(--border-light);
    margin-top: 1rem;
}

/* ===== STATUS INDICATORS ===== */
.status-badge {
    display: inline-flex;
    align-items: center;
    gap: 0.375rem;
    padding: 0.375rem 0.75rem;
    border-radius: 9999px;
    font-size: 0.75rem;
    font-weight: 500;
}

.status-success {
    background: rgba(16, 185, 129, 0.1);
    color: #059669;
}

.status-warning {
    background: rgba(245, 158, 11, 0.1);
    color: #D97706;
}

.status-error {
    background: rgba(239, 68, 68, 0.1);
    color: #DC2626;
}

/* ===== RESPONSIVE ===== */
@media (max-width: 768px) {
    .header-container {
        padding: 1.5rem;
    }
    
    .header-title {
        font-size: 1.5rem;
    }
    
    .progress-stepper {
        flex-wrap: wrap;
        gap: 1rem;
    }
    
    .step {
        flex: 0 0 25%;
    }
    
    .action-bar {
        flex-wrap: wrap;
    }
}

/* ===== MARKDOWN RENDERING ===== */
.markdown-body {
    font-family: inherit !important;
}

.markdown-body table {
    width: 100%;
    border-collapse: collapse;
    margin: 1rem 0;
}

.markdown-body th,
.markdown-body td {
    padding: 0.75rem;
    border: 1px solid var(--border-light);
    text-align: left;
}

.markdown-body th {
    background: var(--bg-primary);
    font-weight: 600;
}

/* Hide default gradio elements */
.gradio-container footer {
    display: none !important;
}
"""

# =============================================================================
# PROGRESS STATE MANAGEMENT
# =============================================================================

PIPELINE_STEPS = [
    {"id": "search", "label": "Searching", "icon": "🔍"},
    {"id": "rank", "label": "Ranking", "icon": "📊"},
    {"id": "download", "label": "Downloading", "icon": "⬇️"},
    {"id": "extract", "label": "Extracting", "icon": "📄"},
    {"id": "analyze", "label": "Analyzing", "icon": "🔬"},
    {"id": "write", "label": "Writing", "icon": "✍️"},
    {"id": "critique", "label": "Critiquing", "icon": "📝"},
    {"id": "finalize", "label": "Complete", "icon": "✅"},
]


def create_progress_html(current_step: int, status: str = "running") -> str:
    """Generate progress stepper HTML."""
    html = '<div class="progress-stepper">'
    
    for idx, step in enumerate(PIPELINE_STEPS):
        if idx < current_step:
            state_class = "step-complete"
            icon = "✓"
        elif idx == current_step:
            state_class = "step-active" if status == "running" else "step-complete"
            icon = "●" if status == "running" else "✓"
        else:
            state_class = "step-pending"
            icon = str(idx + 1)
        
        html += f'''
        <div class="step {state_class}">
            <div class="step-icon">{icon}</div>
            <span class="step-label">{step["label"]}</span>
        </div>
        '''
    
    html += '</div>'
    return html


# =============================================================================
# PIPELINE EXECUTION
# =============================================================================

def run_review_pipeline(
    topic: str,
    keywords: str,
    num_papers: int,
    progress=gr.Progress()
) -> Tuple[Dict, str, str, str, str, str, str, str]:
    """
    Execute the complete literature review pipeline.
    
    Returns:
        Tuple of (status, progress_html, abstract, intro, methods, results, conclusion, references)
    """
    results = {
        "status": "running",
        "current_step": 0,
        "papers_found": 0,
        "papers_selected": 0,
        "errors": []
    }
    
    # Combine topic and keywords
    search_query = topic
    if keywords and keywords.strip():
        search_query += " " + keywords.replace(",", " ")
    
    # Trusted open-access hosts (will be prioritized in download order)
    PREFERRED_HOSTS = (
        "arxiv.org",
        "biorxiv.org",
        "medrxiv.org",
        "pdfs.semanticscholar.org",
        "s2-",
        "semanticscholar.org",
        "openreview.net",
        "core.ac.uk",
        "openaccess.thecvf.com",
        "aclanthology.org",
        "ceur-ws.org",
        "hal.science",
        "eprints",
        "researchgate.net",
    )
    
    # Paywalled hosts to block completely
    BLOCKED_HOSTS = (
        "ieeexplore.ieee.org",
        "dl.acm.org",
        "link.springer.com",
        "wiley.com",
        "sciencedirect.com",
        "onlinelibrary.wiley.com",
        "tandfonline.com",
        "emerald.com",
    )
    
    def get_host(url: str) -> str:
        """Extract host from URL."""
        try:
            return urlparse(url).netloc.lower()
        except Exception:
            return ""
    
    def is_preferred_host(url: str) -> bool:
        """Check if URL is from a preferred open-access host."""
        host = get_host(url)
        return any(ph in host for ph in PREFERRED_HOSTS)
    
    def is_blocked_host(url: str) -> bool:
        """Check if URL is from a blocked paywalled host."""
        host = get_host(url)
        return any(bh in host for bh in BLOCKED_HOSTS)
    
    def is_acceptable_pdf(p: Dict[str, Any]) -> bool:
        """Check if paper has acceptable PDF URL."""
        if not p.get("pdf_available"):
            return False
        url = p.get("pdf_url") or ""
        if not url:
            return False
        return not is_blocked_host(url)
    
    def pdf_priority(p: Dict[str, Any]) -> int:
        """Return priority for sorting (lower = better)."""
        url = p.get("pdf_url") or ""
        if is_preferred_host(url):
            return 0  # Best: known open-access hosts
        elif "arxiv" in url.lower():
            return 0  # Best: arxiv
        elif is_blocked_host(url):
            return 99  # Worst: paywalled
        else:
            return 1  # Medium: unknown hosts
    
    try:
        # Step 1: Search papers - use multiple strategies to find open-access papers
        progress(0.1, desc="🔍 Searching for papers...")
        results["current_step"] = 0
        
        all_papers: List[Dict[str, Any]] = []
        seen_ids: set = set()
        
        # Strategy 1: Direct search with open access terms
        search_results = search_papers(
            topic=f"{search_query} open access",
            limit=30,  # Get many candidates
            year_min=2018,
            min_citations=3  # Lower threshold for more results
        )
        for p in search_results.get("papers", []):
            pid = p.get("paper_id")
            if pid and pid not in seen_ids:
                all_papers.append(p)
                seen_ids.add(pid)
        
        # Strategy 2: Search with arxiv term for guaranteed open access
        progress(0.12, desc="🔍 Searching arxiv...")
        arxiv_results = search_papers(
            topic=f"{search_query} arxiv",
            limit=20,
            year_min=2018,
            min_citations=0  # No citation requirement for arxiv
        )
        for p in arxiv_results.get("papers", []):
            pid = p.get("paper_id")
            if pid and pid not in seen_ids:
                all_papers.append(p)
                seen_ids.add(pid)
        
        # Strategy 3: Plain search as fallback
        if len(all_papers) < 15:
            progress(0.14, desc="🔍 Fallback search...")
            fallback_results = search_papers(
                topic=search_query,
                limit=25,
                year_min=2015,
                min_citations=3
            )
            for p in fallback_results.get("papers", []):
                pid = p.get("paper_id")
                if pid and pid not in seen_ids:
                    all_papers.append(p)
                    seen_ids.add(pid)
        
        results["papers_found"] = len(all_papers)
        
        if not all_papers:
            return (
                {"status": "error", "message": "No papers found"},
                create_progress_html(0, "error"),
                "No papers found for this topic.",
                "", "", "", "", ""
            )
        
        # Step 2: Rank and filter papers
        progress(0.2, desc="📊 Ranking papers...")
        results["current_step"] = 1
        
        # First, filter for acceptable PDFs only
        papers_with_pdf = [p for p in all_papers if is_acceptable_pdf(p)]
        
        # Sort by PDF host preference (arxiv/open-access first)
        papers_with_pdf.sort(key=pdf_priority)
        
        # Now rank by relevance
        ranked_papers = rank_papers(papers_with_pdf, search_query, top_n=min(20, len(papers_with_pdf)))
        
        # Re-sort to prefer open-access hosts among top ranked
        def combined_score(p: Dict[str, Any]) -> tuple:
            """Sort by (host_priority, -rank)."""
            return (pdf_priority(p), -(p.get("rank", 0)))
        
        ranked_papers.sort(key=combined_score)
        
        # Take top candidates for download attempts
        pdf_candidates = ranked_papers[:15] if len(ranked_papers) > 5 else ranked_papers
        
        if not pdf_candidates:
            return (
                {"status": "error", "message": "No papers with PDFs available"},
                create_progress_html(1, "error"),
                "No papers with downloadable PDFs found.",
                "", "", "", "", ""
            )
        
        # Save metadata for later reference
        metadata_to_save = {
            "topic": search_query,
            "timestamp": datetime.now().isoformat(),
            "papers_found": len(all_papers),
            "papers_with_pdf": len(papers_with_pdf),
            "papers": ranked_papers
        }
        save_metadata(metadata_to_save)
        
        # Step 3: Download PDFs
        progress(0.35, desc="⬇️ Downloading PDFs...")
        results["current_step"] = 2
        
        downloaded_papers = download_papers(pdf_candidates)
        # Keep only top 3 successful downloads (prefer open-access sources)
        downloaded_papers.sort(key=pdf_priority)
        downloaded_papers = downloaded_papers[:3]
        results["papers_selected"] = len(downloaded_papers)
        
        # If fewer than 3, run a supplemental open-access search
        if len(downloaded_papers) < 3:
            try:
                progress(0.4, desc="🔎 Supplemental arxiv search…")
                # Try multiple supplemental searches
                supplemental_queries = [
                    f"{search_query} arxiv preprint",
                    f"{search_query} open access pdf",
                ]
                
                for sup_query in supplemental_queries:
                    if len(downloaded_papers) >= 3:
                        break
                        
                    supplemental = search_papers(
                        topic=sup_query,
                        limit=25,
                        year_min=2015,
                        min_citations=0  # No citation filter for supplemental
                    )
                    sup_papers = supplemental.get("papers", [])
                    if sup_papers:
                        # Filter and prioritize
                        sup_with_pdf = [p for p in sup_papers if is_acceptable_pdf(p)]
                        sup_with_pdf.sort(key=pdf_priority)
                        
                        # Avoid duplicates
                        existing_ids = {p.get("paper_id") for p in downloaded_papers}
                        sup_candidates = [p for p in sup_with_pdf if p.get("paper_id") not in existing_ids]
                        
                        if sup_candidates:
                            # Download up to what we need
                            needed = 3 - len(downloaded_papers)
                            more = download_papers(sup_candidates[:needed * 3])
                            if more:
                                downloaded_papers.extend(more)
                                downloaded_papers = downloaded_papers[:3]
                                results["papers_selected"] = len(downloaded_papers)
            except Exception as sup_err:
                logger.warning(f"Supplemental search error: {sup_err}")

        if not downloaded_papers:
            return (
                {"status": "partial", "message": "PDF download failed"},
                create_progress_html(2, "error"),
                "Could not download papers.", "", "", "", "", ""
            )
        
        # Step 4: Extract text
        progress(0.5, desc="📄 Extracting text...")
        results["current_step"] = 3
        
        downloaded_papers = extract_text_for_papers(downloaded_papers)
        save_download_metadata(downloaded_papers, "selected_papers.json")
        
        # Step 5: Analyze
        progress(0.65, desc="🔬 Analyzing papers...")
        results["current_step"] = 4
        
        analysis = process_analysis()
        
        # Step 6: Generate draft
        progress(0.8, desc="✍️ Writing literature review...")
        results["current_step"] = 5
        
        drafts = generate_drafts()
        review = drafts.get("literature_review", {})
        sections = review.get("sections", {})
        
        # Step 7: Critique
        progress(0.9, desc="📝 Critiquing draft...")
        results["current_step"] = 6
        
        critique_result = critique_drafts()
        
        # Step 8: Complete
        progress(1.0, desc="✅ Complete!")
        results["current_step"] = 7
        results["status"] = "success"
        
        # Extract sections
        abstract = sections.get("abstract", "Abstract not generated.")
        introduction = sections.get("introduction", "Introduction not generated.")
        methods = sections.get("methodology_comparison", "Methods section not generated.")
        results_section = sections.get("results_synthesis", "Results section not generated.")
        conclusion = sections.get("conclusion", "Conclusion not generated.")
        references = sections.get("references", "References not generated.")
        
        return (
            results,
            create_progress_html(7, "complete"),
            abstract,
            introduction,
            methods,
            results_section,
            conclusion,
            references
        )
        
    except Exception as e:
        import traceback
        error_msg = str(e)
        results["status"] = "error"
        results["errors"].append(error_msg)
        
        return (
            results,
            create_progress_html(results["current_step"], "error"),
            f"Error: {error_msg}",
            "", "", "", "", ""
        )


def run_critique_revision() -> Tuple[str, str]:
    """Run critique and return suggestions with revised content."""
    try:
        critique_result = critique_drafts()
        critique = critique_result.get("critique", {})
        revised = critique_result.get("revised_review", {})
        
        # Format critique output
        output = f"""## Quality Assessment

**Overall Score:** {critique.get('overall_score', 0):.0%} ({critique.get('grade', 'N/A')})

**Verdict:** {critique.get('verdict', 'No verdict available')}

### Suggestions for Improvement

"""
        for idx, suggestion in enumerate(critique.get("all_suggestions", [])[:10], 1):
            output += f"{idx}. {suggestion}\n"
        
        revised_text = revised.get("full_text", "No revision generated.")
        
        return output, revised_text
    except Exception as e:
        return f"Error during critique: {str(e)}", ""


# =============================================================================
# PDF EXPORT (SERVER-SIDE)
# =============================================================================

def generate_pdf_export() -> Optional[str]:
    """
    Generate PDF export using ReportLab.
    Returns path to generated PDF file.
    """
    try:
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.lib.colors import HexColor
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, PageBreak,
            Table, TableStyle, ListFlowable, ListItem
        )
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
        
        # Load draft data
        drafts_path = ROOT / "data" / "metadata" / "drafts.json"
        if not drafts_path.exists():
            return None
        
        with open(drafts_path, "r", encoding="utf-8") as f:
            drafts_data = json.load(f)
        
        review = drafts_data.get("literature_review", {})
        sections = review.get("sections", {})
        topic = review.get("topic", "Literature Review")
        
        # Create PDF
        output_path = ROOT / "data" / "literature_review_export.pdf"
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=1*inch,
            leftMargin=1*inch,
            topMargin=1*inch,
            bottomMargin=1*inch
        )
        
        # Define styles
        styles = getSampleStyleSheet()
        
        # Title style
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=24,
            spaceAfter=30,
            alignment=TA_CENTER,
            textColor=HexColor('#1F2937'),
            fontName='Helvetica-Bold'
        )
        
        # Subtitle style
        subtitle_style = ParagraphStyle(
            'Subtitle',
            parent=styles['Normal'],
            fontSize=12,
            spaceAfter=40,
            alignment=TA_CENTER,
            textColor=HexColor('#6B7280')
        )
        
        # Heading style
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading1'],
            fontSize=16,
            spaceBefore=24,
            spaceAfter=12,
            textColor=HexColor('#1F2937'),
            fontName='Helvetica-Bold',
            borderPadding=(0, 0, 4, 0),
            borderColor=HexColor('#E5E7EB'),
            borderWidth=0
        )
        
        # Body style
        body_style = ParagraphStyle(
            'CustomBody',
            parent=styles['Normal'],
            fontSize=11,
            spaceAfter=12,
            alignment=TA_JUSTIFY,
            leading=16,
            textColor=HexColor('#374151')
        )
        
        # Reference style
        ref_style = ParagraphStyle(
            'Reference',
            parent=styles['Normal'],
            fontSize=10,
            spaceAfter=8,
            leftIndent=36,
            firstLineIndent=-36,
            textColor=HexColor('#4B5563')
        )
        
        # Build content
        story = []
        
        # Title page
        story.append(Spacer(1, 2*inch))
        story.append(Paragraph("Literature Review", title_style))
        story.append(Paragraph(topic.title(), subtitle_style))
        story.append(Spacer(1, 0.5*inch))
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y')}", subtitle_style))
        story.append(Paragraph(f"Papers Reviewed: {review.get('papers_count', 0)}", subtitle_style))
        story.append(PageBreak())
        
        # Section order
        section_config = [
            ("abstract", "Abstract"),
            ("introduction", "Introduction"),
            ("methodology_comparison", "Methods"),
            ("results_synthesis", "Results"),
            ("conclusion", "Conclusion"),
            ("references", "References")
        ]
        
        for section_key, section_title in section_config:
            content = sections.get(section_key, "")
            if not content:
                continue
            
            # Section heading
            story.append(Paragraph(section_title, heading_style))
            
            # Clean markdown
            clean_content = content
            clean_content = clean_content.replace("## ", "").replace("### ", "")
            clean_content = clean_content.replace("**", "").replace("*", "")
            clean_content = clean_content.replace("|", " ")
            
            # Handle references specially
            if section_key == "references":
                for line in clean_content.split("\n"):
                    line = line.strip()
                    if line and not line.startswith("-"):
                        # Escape special characters
                        safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                        story.append(Paragraph(safe_line, ref_style))
            else:
                # Regular paragraphs
                paragraphs = clean_content.split("\n\n")
                for para in paragraphs:
                    para = para.strip()
                    if para and not para.startswith("|") and not para.startswith("-"):
                        # Escape special characters
                        safe_para = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                        story.append(Paragraph(safe_para, body_style))
        
        # Build PDF
        doc.build(story)
        
        return str(output_path)
        
    except ImportError:
        return None
    except Exception as e:
        print(f"PDF generation error: {e}")
        return None


def generate_markdown_export() -> Optional[str]:
    """Generate Markdown export."""
    try:
        drafts_path = ROOT / "data" / "metadata" / "drafts.json"
        if not drafts_path.exists():
            return None
        
        with open(drafts_path, "r", encoding="utf-8") as f:
            drafts_data = json.load(f)
        
        review = drafts_data.get("literature_review", {})
        full_text = review.get("full_text", "")
        
        if not full_text:
            return None
        
        output_path = ROOT / "data" / "literature_review_export.md"
        
        # Add header
        header = f"""# Literature Review: {review.get('topic', 'Research Topic')}

**Generated:** {datetime.now().strftime('%B %d, %Y')}  
**Papers Reviewed:** {review.get('papers_count', 0)}

---

"""
        
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(header + full_text)
        
        return str(output_path)
        
    except Exception as e:
        print(f"Markdown export error: {e}")
        return None


def generate_docx_export() -> Optional[str]:
    """Generate DOCX export."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
        
        drafts_path = ROOT / "data" / "metadata" / "drafts.json"
        if not drafts_path.exists():
            return None
        
        with open(drafts_path, "r", encoding="utf-8") as f:
            drafts_data = json.load(f)
        
        review = drafts_data.get("literature_review", {})
        sections = review.get("sections", {})
        topic = review.get("topic", "Literature Review")
        
        # Create document
        doc = Document()
        
        # Title
        title = doc.add_heading(f"Literature Review: {topic.title()}", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Metadata
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y')}\n")
        meta.add_run(f"Papers Reviewed: {review.get('papers_count', 0)}")
        
        doc.add_paragraph()  # Spacer
        
        # Sections
        section_config = [
            ("abstract", "Abstract"),
            ("introduction", "Introduction"),
            ("methodology_comparison", "Methods"),
            ("results_synthesis", "Results"),
            ("conclusion", "Conclusion"),
            ("references", "References")
        ]
        
        for section_key, section_title in section_config:
            content = sections.get(section_key, "")
            if not content:
                continue
            
            doc.add_heading(section_title, 1)
            
            # Clean and add content
            clean_content = content.replace("## ", "").replace("### ", "")
            clean_content = clean_content.replace("**", "").replace("*", "")
            clean_content = clean_content.replace("|", " ")
            
            for para in clean_content.split("\n\n"):
                para = para.strip()
                if para and not para.startswith("|") and not para.startswith("-"):
                    doc.add_paragraph(para)
        
        output_path = ROOT / "data" / "literature_review_export.docx"
        doc.save(str(output_path))
        
        return str(output_path)
        
    except ImportError:
        return None
    except Exception as e:
        print(f"DOCX export error: {e}")
        return None


def export_pdf():
    """Export handler for PDF button."""
    path = generate_pdf_export()
    if path and os.path.exists(path):
        return gr.File(value=path, visible=True)
    return gr.File(value=None, visible=False)


def export_markdown():
    """Export handler for Markdown button."""
    path = generate_markdown_export()
    if path and os.path.exists(path):
        return gr.File(value=path, visible=True)
    return gr.File(value=None, visible=False)


def export_docx():
    """Export handler for DOCX button."""
    path = generate_docx_export()
    if path and os.path.exists(path):
        return gr.File(value=path, visible=True)
    return gr.File(value=None, visible=False)


# =============================================================================
# GRADIO INTERFACE
# =============================================================================

def create_app():
    """Create and configure the Gradio application."""
    
    with gr.Blocks(
        title="AI Literature Review System"
    ) as app:
        
        # ===== HEADER =====
        gr.HTML("""
        <div class="header-container">
            <h1 class="header-title">🎓 AI Literature Review System</h1>
            <p class="header-subtitle">
                Automated systematic review using Semantic Scholar and intelligent analysis
            </p>
        </div>
        """)
        
        # ===== INPUT SECTION =====
        with gr.Group(elem_classes=["card"]):
            gr.HTML('<div class="card-title">📝 Research Parameters</div>')
            
            with gr.Row():
                with gr.Column(scale=2):
                    topic_input = gr.Textbox(
                        label="Research Topic",
                        placeholder="e.g., deep learning for medical image analysis",
                        max_lines=1,
                        elem_classes=["input-container"]
                    )
                with gr.Column(scale=2):
                    keywords_input = gr.Textbox(
                        label="Additional Keywords (optional)",
                        placeholder="e.g., CNN, segmentation, diagnosis",
                        max_lines=1,
                        elem_classes=["input-container"]
                    )
            
            with gr.Row():
                with gr.Column(scale=2):
                    papers_slider = gr.Slider(
                        minimum=10,
                        maximum=20,
                        value=15,
                        step=1,
                        label="Number of Papers to Search"
                    )
                with gr.Column(scale=1):
                    run_button = gr.Button(
                        "▶️ Run Review",
                        variant="primary",
                        elem_classes=["primary-button"]
                    )
        
        # ===== PROGRESS SECTION =====
        with gr.Group(elem_classes=["card"]):
            gr.HTML('<div class="card-title">📊 Progress</div>')
            progress_display = gr.HTML(
                value=create_progress_html(-1, "pending"),
                elem_id="progress-stepper"
            )
            status_output = gr.JSON(visible=False)
        
        # ===== RESULTS SECTION =====
        with gr.Group(elem_classes=["card"]):
            gr.HTML('<div class="card-title">📚 Literature Review</div>')
            
            with gr.Tabs(elem_classes=["tabs-container"]) as result_tabs:
                with gr.TabItem("Abstract", id="tab-abstract"):
                    abstract_output = gr.Markdown(
                        value="*Run a review to generate content...*",
                        elem_classes=["content-area"]
                    )
                
                with gr.TabItem("Introduction", id="tab-intro"):
                    intro_output = gr.Markdown(
                        value="*Run a review to generate content...*",
                        elem_classes=["content-area"]
                    )
                
                with gr.TabItem("Methods", id="tab-methods"):
                    methods_output = gr.Markdown(
                        value="*Run a review to generate content...*",
                        elem_classes=["content-area"]
                    )
                
                with gr.TabItem("Results", id="tab-results"):
                    results_output = gr.Markdown(
                        value="*Run a review to generate content...*",
                        elem_classes=["content-area"]
                    )
                
                with gr.TabItem("Conclusion", id="tab-conclusion"):
                    conclusion_output = gr.Markdown(
                        value="*Run a review to generate content...*",
                        elem_classes=["content-area"]
                    )
                
                with gr.TabItem("References", id="tab-refs"):
                    refs_output = gr.Markdown(
                        value="*Run a review to generate content...*",
                        elem_classes=["content-area"]
                    )
        
        # ===== ACTIONS SECTION =====
        with gr.Group(elem_classes=["card"]):
            gr.HTML('<div class="card-title">📤 Actions</div>')
            
            with gr.Row():
                with gr.Column(scale=1):
                    critique_btn = gr.Button(
                        "🔄 Critique / Revise",
                        variant="secondary",
                        elem_classes=["secondary-button"]
                    )
                with gr.Column(scale=1):
                    pdf_btn = gr.Button(
                        "📄 Download PDF",
                        variant="secondary",
                        elem_classes=["secondary-button"]
                    )
                with gr.Column(scale=1):
                    md_btn = gr.Button(
                        "📝 Download Markdown",
                        variant="secondary",
                        elem_classes=["secondary-button"]
                    )
                with gr.Column(scale=1):
                    docx_btn = gr.Button(
                        "📑 Download DOCX",
                        variant="secondary",
                        elem_classes=["secondary-button"]
                    )
            
            # Hidden file outputs for downloads
            pdf_file = gr.File(visible=False, label="PDF Download")
            md_file = gr.File(visible=False, label="Markdown Download")
            docx_file = gr.File(visible=False, label="DOCX Download")
            
            # Critique output
            with gr.Accordion("Critique Results", open=False):
                critique_output = gr.Markdown(value="")
                revised_output = gr.Markdown(value="")
        
        # ===== EVENT HANDLERS =====
        
        # Run pipeline
        run_button.click(
            fn=run_review_pipeline,
            inputs=[topic_input, keywords_input, papers_slider],
            outputs=[
                status_output,
                progress_display,
                abstract_output,
                intro_output,
                methods_output,
                results_output,
                conclusion_output,
                refs_output
            ],
            show_progress="full"
        )
        
        # Critique/Revise
        critique_btn.click(
            fn=run_critique_revision,
            inputs=[],
            outputs=[critique_output, revised_output]
        )
        
        # Export handlers
        pdf_btn.click(fn=export_pdf, inputs=[], outputs=[pdf_file])
        md_btn.click(fn=export_markdown, inputs=[], outputs=[md_file])
        docx_btn.click(fn=export_docx, inputs=[], outputs=[docx_file])
    
    return app


# =============================================================================
# MAIN
# =============================================================================

if __name__ == "__main__":
    print("\n" + "=" * 60)
    print("🎓 AI Literature Review System")
    print("=" * 60)
    print("\n📍 Starting server...")
    print("🌐 Open: http://127.0.0.1:7860")
    print("\n💡 Press Ctrl+C to stop\n")
    
    app = create_app()
    app.launch(
        server_name="127.0.0.1",
        server_port=7860,
        share=False,
        show_error=True,
        css=CUSTOM_CSS,
        theme=gr.themes.Soft(
            primary_hue="blue",
            secondary_hue="slate",
            neutral_hue="slate"
        )
    )
